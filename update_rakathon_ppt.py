#!/usr/bin/env python3
"""Update Rakathon ASAKAI deck per editorial brief."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PPT_PATH = "/workspace/Rakathon_2026_ASAKAI.pptx"
FONT = "Rakuten Sans JP"
RED = RGBColor(0xBF, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_MID = RGBColor(0x9C, 0xA3, 0xAF)
GRAY_BODY = RGBColor(0xD1, 0xD5, 0xDB)
GRAY_EYE = RGBColor(0x6B, 0x72, 0x80)


def style_run(run, *, size_pt, color, bold=False, spacing_pt=None):
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    if spacing_pt is not None:
        run.font.character_spacing = Pt(spacing_pt)


def add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def set_cover_slide(slide):
    """Mirror slide 3 title hierarchy on black cover (slide-master layout)."""
    # Red accent bar (aligned with slide 3 PROGRAMME TIMELINE block)
    bar = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        Inches(0.75),
        Inches(0.77),
        Inches(0.045),
        Inches(0.7),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED
    bar.line.fill.background()

    eyebrow = add_textbox(slide, Inches(0.88), Inches(0.55), Inches(8.5), Inches(0.25))
    tf = eyebrow.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = "ASAKAI · "
    style_run(r1, size_pt=9, color=GRAY_EYE, spacing_pt=2)
    r2 = p.add_run()
    r2.text = "RAKATHON 2026"
    style_run(r2, size_pt=9, color=WHITE, spacing_pt=2)

    title_box = add_textbox(slide, Inches(0.88), Inches(0.81), Inches(11.5), Inches(0.85))
    tp = title_box.text_frame.paragraphs[0]
    rt1 = tp.add_run()
    rt1.text = "Rakathon"
    style_run(rt1, size_pt=44, color=WHITE)
    rt2 = tp.add_run()
    rt2.text = ": From Idea to D-Day"
    style_run(rt2, size_pt=44, color=WHITE)

    sub_box = add_textbox(slide, Inches(0.88), Inches(1.63), Inches(11.0), Inches(0.6))
    stf = sub_box.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sr = sp.add_run()
    sr.text = (
        "Submit your idea, get shortlisted, hack for a week, "
        "and demo live to the senior jury"
    )
    style_run(sr, size_pt=14, color=GRAY_BODY)

    meta_box = add_textbox(slide, Inches(0.88), Inches(6.55), Inches(10.0), Inches(0.35))
    mp = meta_box.text_frame.paragraphs[0]
    mr = mp.add_run()
    mr.text = "Submission 15–31 Oct 2026  ·  Event Day 23 Nov 2026  ·  RIEPL India"
    style_run(mr, size_pt=10, color=GRAY_MID, spacing_pt=1)


def content_placeholder(slide):
    for shape in slide.placeholders:
        if shape.is_placeholder and shape.placeholder_format.idx == 10:
            return shape
    raise RuntimeError("content placeholder not found")


def replace_body_text(shape, lines, *, bullet=True):
    tf = shape.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        if bullet:
            p.font.name = FONT
            p.font.size = Pt(18)


def set_slide4(slide):
    slide.shapes.title.text = "Rakathon 2026 — Background"
    body = content_placeholder(slide)
    tf = body.text_frame
    tf.clear()

    p0 = tf.paragraphs[0]
    p0.text = "Theme: Enterprise Agentic AI — build Super Agents that create real business impact"
    p0.level = 0

    p1 = tf.add_paragraph()
    p1.text = "2025 Hack-a-Prompt (vibe coding) → 2026 Rakathon: production-ready agentic solutions"
    p1.level = 0

    p2 = tf.add_paragraph()
    p2.text = "Aligned with RPC 2026 and Mikitani-san's Super Agent vision"
    p2.level = 0

    p3 = tf.add_paragraph()
    p3.text = "Tagline contest coming soon — help us name Rakathon 2026"
    p3.level = 0

    p4 = tf.add_paragraph()
    p4.level = 0
    run = p4.add_run()
    run.text = "Reference: "
    run.font.name = FONT
    run.font.size = Pt(16)
    link = p4.add_run()
    link.text = "Super Agent at Rakuten AI Optimism 2026"
    link.font.name = FONT
    link.font.size = Pt(16)
    link.font.color.rgb = RED
    link.hyperlink.address = (
        "https://rakuteninternational.com/news/"
        "mikitani-unveils-the-super-agent-at-rakuten-ai-optimism-2026"
    )

    for para in tf.paragraphs:
        para.font.name = FONT
        if para.font.size is None:
            para.font.size = Pt(18)


def set_slide5(slide):
    slide.shapes.title.text = "What we need from you"
    replace_body_text(
        content_placeholder(slide),
        [
            "Tagline contest — submit a creative tagline for Rakathon 2026",
            "Idea submission — 1-page proposal, POC, and expected business impact (by 31 Oct)",
            "Problem areas — prepare your pitch around leader-identified RIEPL focus areas",
            "Teams — 1 to 4 members; cross-department collaboration encouraged",
            "Recruit teammates and register during the submission window",
        ],
    )


def set_slide6(slide):
    slide.shapes.title.text = "What you get out of the event"
    replace_body_text(
        content_placeholder(slide),
        [
            "Chance to attend a prominent AI conference — and more for top teams",
            "India Agent Marketplace launch — host your agents for use across RIEPL",
            "Impact on leader-identified problem areas with org-wide visibility",
            "Mentorship during Hack Week and recognition for Top 3 winners",
        ],
    )


def main():
    prs = Presentation(PPT_PATH)
    set_cover_slide(prs.slides[0])
    set_slide4(prs.slides[3])
    set_slide5(prs.slides[4])
    set_slide6(prs.slides[5])
    prs.save(PPT_PATH)
    print("Saved", PPT_PATH)


if __name__ == "__main__":
    main()
